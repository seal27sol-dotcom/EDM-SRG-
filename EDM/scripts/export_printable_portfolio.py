from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def add_markdown_like_text(document: Document, markdown_text: str) -> None:
    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            document.add_paragraph()
            continue

        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("**") and line.endswith("**"):
            document.add_paragraph(line[2:-2], style="Intense Quote")
        elif line.startswith("|"):
            document.add_paragraph(line)
        elif line.startswith("```"):
            continue
        else:
            document.add_paragraph(line)


def build_pdf(markdown_text: str, out_path: Path) -> None:
    styles = getSampleStyleSheet()
    style_body = styles["BodyText"]
    style_body.fontName = "Helvetica"
    style_body.fontSize = 10
    style_body.leading = 13
    style_title = styles["Title"]
    style_title.fontName = "Helvetica-Bold"
    style_h2 = styles["Heading2"]
    style_h2.fontName = "Helvetica-Bold"
    style_h3 = styles["Heading3"]
    style_h3.fontName = "Helvetica-Bold"

    story = []
    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            story.append(Spacer(1, 8))
            continue
        if line.startswith("# "):
            story.append(Paragraph(line[2:], style_title))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], style_h2))
        elif line.startswith("### "):
            story.append(Paragraph(line[4:], style_h3))
        elif line.startswith("- "):
            story.append(Paragraph(f"• {line[2:]}", style_body))
        elif line.startswith("|"):
            story.append(Paragraph(line, style_body))
        elif line.startswith("```"):
            continue
        else:
            story.append(Paragraph(line, style_body))

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=A4,
        leftMargin=25 * mm,
        rightMargin=25 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
    )
    doc.build(story)


def main() -> None:
    base = Path(__file__).resolve().parents[1]
    source = base / 'docs' / 'portfolio.md'
    docs_dir = source.parent
    markdown_text = source.read_text(encoding='utf-8')

    docx_path = docs_dir / 'portfolio_print.docx'
    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    add_markdown_like_text(doc, markdown_text)
    doc.save(docx_path)

    pdf_path = docs_dir / 'portfolio_print.pdf'
    build_pdf(markdown_text, pdf_path)

    print(f'Created DOCX: {docx_path}')
    print(f'Created PDF: {pdf_path}')


if __name__ == '__main__':
    from docx.shared import Pt
    main()
